"""客服知识 Excel 导入 / 导出 / 模板。

从 UI 里抽出来单独成模块，原因有两个：一是解析逻辑要能单测，
二是导出与导入必须共用同一套列定义，否则「导出改完再导回去」会对不上。

不依赖 PyQt，也不直接写数据库：解析只负责把表格变成行数据，
入库仍交给 KnowledgeService.batch_import_customer_service。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

# 标准列：导出与模板都用这一套，导入优先按这些名字识别。
COLUMN_TITLE = "标题"
COLUMN_CONTENT = "内容"
COLUMN_TAGS = "标签"
COLUMN_ENABLED = "启用"
STANDARD_HEADERS: Tuple[str, ...] = (
    COLUMN_TITLE,
    COLUMN_CONTENT,
    COLUMN_TAGS,
    COLUMN_ENABLED,
)

# 表头别名 -> 内部字段。键统一小写去空格后比较。
# 「一级分类 / 二级分类」是本项目历史导入格式，必须继续认，
# 否则用户手上已经在用的 Excel 会突然失效。
_HEADER_ALIASES: Dict[str, str] = {
    "标题": "title",
    "话术标题": "title",
    "问题": "title",
    "title": "title",
    "内容": "content",
    "话术内容": "content",
    "答案": "content",
    "回复": "content",
    "content": "content",
    "标签": "tags",
    "分类": "tags",
    "tags": "tags",
    "tag": "tags",
    "启用": "enabled",
    "状态": "enabled",
    "是否启用": "enabled",
    "enabled": "enabled",
    "一级分类": "category1",
    "二级分类": "category2",
}

# 表头完全认不出来时的兜底列序，等同于历史实现的硬编码顺序。
_LEGACY_POSITIONS: Tuple[str, ...] = ("category1", "category2", "title", "content")

_TRUE_WORDS = frozenset({"1", "true", "yes", "y", "t", "是", "启用", "开启", "有效", "正常", "on"})
_FALSE_WORDS = frozenset({"0", "false", "no", "n", "f", "否", "禁用", "停用", "关闭", "无效", "off"})

# Excel 会把 = + - @ 开头的文本当公式执行，导出前必须锁成纯文本。
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# 表头占 1 行，pandas 的第 0 行数据对应 Excel 第 2 行。
_FIRST_DATA_ROW = 2


@dataclass
class ParseResult:
    """解析结果。

    rows 直接喂给 batch_import_customer_service；skipped 用于向用户交代
    每一条为什么没进来——只给一个总数的话，用户无从修起。
    """

    rows: List[Dict[str, Any]] = field(default_factory=list)
    skipped: List[Tuple[int, str]] = field(default_factory=list)
    layout: str = "standard"
    header_recognized: bool = True

    @property
    def skipped_count(self) -> int:
        return len(self.skipped)

    def describe_skipped(self, limit: int = 50) -> str:
        """给 UI 用的多行说明，超出上限只提示剩余条数，避免弹窗撑爆。"""
        if not self.skipped:
            return ""
        lines = [f"第 {row} 行：{reason}" for row, reason in self.skipped[:limit]]
        remaining = len(self.skipped) - limit
        if remaining > 0:
            lines.append(f"…… 另有 {remaining} 行被跳过")
        return "\n".join(lines)


def _normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    # pandas 遇到重复列名会加 .1/.2 后缀，去掉后再比对别名。
    text = text.strip().lower().replace(" ", "").replace("　", "")
    if "." in text:
        head, _, tail = text.rpartition(".")
        if head and tail.isdigit():
            text = head
    return text


def _map_headers(columns: Sequence[Any]) -> Dict[str, int]:
    """表头名 -> 列下标。同一字段出现多次时以第一次为准。"""
    mapping: Dict[str, int] = {}
    for index, raw in enumerate(columns):
        field_name = _HEADER_ALIASES.get(_normalize_header(raw))
        if field_name and field_name not in mapping:
            mapping[field_name] = index
    return mapping


def parse_enabled(value: Any, default: bool = True) -> Optional[bool]:
    """解析启用列。空值取默认；无法识别返回 None，由调用方决定是否跳过。"""
    if value is None:
        return default
    text = str(value).strip().lower()
    if not text or text in {"nan", "none"}:
        return default
    if text in _TRUE_WORDS:
        return True
    if text in _FALSE_WORDS:
        return False
    return None


def _cell(values: Sequence[Any], index: Optional[int]) -> str:
    if index is None or index >= len(values):
        return ""
    value = values[index]
    if value is None:
        return ""
    text = str(value).strip()
    # pandas 用 NaN 表示空单元格，dtype=str 下会变成字面量 "nan"。
    return "" if text.lower() in {"nan", "nat", "none"} else text


def _compose_tags(explicit: str, category1: str, category2: str) -> str:
    """标签列优先；没有就用历史的两级分类拼。"""
    if explicit:
        parts = [p.strip() for p in explicit.replace("，", ",").split(",")]
    else:
        parts = [category1, category2]
    seen: List[str] = []
    for part in parts:
        part = part.strip()
        if part and part not in seen:
            seen.append(part)
    return ",".join(seen)


def parse_rows(columns: Sequence[Any], records: Iterable[Sequence[Any]]) -> ParseResult:
    """纯数据解析，不碰文件，方便直接单测。"""
    mapping = _map_headers(columns)
    recognized = bool({"title", "content"} & set(mapping))
    if not recognized:
        # 认不出表头就按历史列序解析，保持旧文件可用；调用方会提示用户。
        mapping = {name: i for i, name in enumerate(_LEGACY_POSITIONS)}

    layout = "standard" if "tags" in mapping or "enabled" in mapping else "legacy"
    result = ParseResult(layout=layout, header_recognized=recognized)

    for offset, record in enumerate(records):
        row_number = _FIRST_DATA_ROW + offset
        values = list(record)
        title = _cell(values, mapping.get("title"))
        content = _cell(values, mapping.get("content"))
        category1 = _cell(values, mapping.get("category1"))
        category2 = _cell(values, mapping.get("category2"))
        tags = _compose_tags(_cell(values, mapping.get("tags")), category1, category2)

        # 整行空白是表格尾部的常见情况，静默跳过更合理。
        if not any([title, content, tags]):
            continue

        missing = []
        if not title:
            missing.append(COLUMN_TITLE)
        if not content:
            missing.append(COLUMN_CONTENT)
        if missing:
            result.skipped.append((row_number, "缺少" + "、".join(missing)))
            continue

        enabled = parse_enabled(values[mapping["enabled"]] if "enabled" in mapping and mapping["enabled"] < len(values) else None)
        if enabled is None:
            raw = _cell(values, mapping.get("enabled"))
            result.skipped.append((row_number, f"启用列无法识别：{raw}"))
            continue

        result.rows.append(
            {"title": title, "content": content, "tags": tags or None, "enabled": enabled}
        )

    return result


def parse_workbook(filepath: str) -> ParseResult:
    """读取 .xlsx / .xls。用 pandas 以兼容两种格式（引擎由它按扩展名选）。"""
    import pandas as pd

    frame = pd.read_excel(filepath, header=0, dtype=str)
    frame = frame.where(pd.notna(frame), None)
    return parse_rows(list(frame.columns), frame.itertuples(index=False, name=None))


def _write_text(worksheet, row: int, column: int, text: str):
    """写纯文本单元格，避免内容被 Excel 当成公式执行。"""
    cell = worksheet.cell(row=row, column=column, value=text)
    if isinstance(text, str) and text[:1] in _FORMULA_PREFIXES:
        cell.data_type = "s"
    return cell


def _style_header(worksheet, headers: Sequence[str], widths: Sequence[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="DDEBF7")
    for index, name in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=name)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = widths[index - 1]
    worksheet.freeze_panes = "A2"


def export_workbook(filepath: str, items: Iterable[Any]) -> int:
    """导出客服知识，返回写出条数。列与模板一致，改完可直接导回。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客服知识"
    _style_header(sheet, STANDARD_HEADERS, (28, 70, 22, 10))

    count = 0
    for offset, item in enumerate(items):
        row = 2 + offset
        tags = getattr(item, "tags", None) or ""
        _write_text(sheet, row, 1, str(getattr(item, "title", "") or ""))
        content_cell = _write_text(sheet, row, 2, str(getattr(item, "content", "") or ""))
        content_cell.alignment = Alignment(vertical="top", wrap_text=True)
        _write_text(sheet, row, 3, str(tags))
        _write_text(sheet, row, 4, "是" if getattr(item, "enabled", True) else "否")
        count += 1

    workbook.save(filepath)
    return count


_TEMPLATE_SAMPLES: Tuple[Tuple[str, str, str, str], ...] = (
    ("发货时间", "亲，我们是当天 16:00 前的订单当天发出，之后的顺延到次日。", "物流,发货", "是"),
    ("七天无理由", "支持七天无理由退换，商品需不影响二次销售，运费由买家承担。", "售后", "是"),
)

_TEMPLATE_NOTES: Tuple[Tuple[str, str], ...] = (
    (COLUMN_TITLE, "必填。一条知识的简短标题，例如「发货时间」。"),
    (COLUMN_CONTENT, "必填。客服实际要用的完整话术内容。"),
    (COLUMN_TAGS, "选填。多个标签用逗号分隔，例如「物流,发货」。"),
    (COLUMN_ENABLED, "选填。填「是」或「否」，留空默认启用。"),
)


def write_template(filepath: str) -> None:
    """生成导入模板：首个工作表可直接被导入器解析，说明单独放一页。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客服知识"
    _style_header(sheet, STANDARD_HEADERS, (28, 70, 22, 10))

    for offset, sample in enumerate(_TEMPLATE_SAMPLES):
        for column, value in enumerate(sample, start=1):
            cell = _write_text(sheet, 2 + offset, column, value)
            if column == 2:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    notes = workbook.create_sheet("填写说明")
    notes.column_dimensions["A"].width = 14
    notes.column_dimensions["B"].width = 78
    title_cell = notes.cell(row=1, column=1, value="填写说明")
    title_cell.font = Font(bold=True, size=12)
    notes.cell(row=2, column=1, value="列名")
    notes.cell(row=2, column=2, value="说明")
    for cell in (notes.cell(row=2, column=1), notes.cell(row=2, column=2)):
        cell.font = Font(bold=True)
    for offset, (name, description) in enumerate(_TEMPLATE_NOTES):
        notes.cell(row=3 + offset, column=1, value=name)
        notes.cell(row=3 + offset, column=2, value=description)

    tail = 3 + len(_TEMPLATE_NOTES) + 1
    for offset, line in enumerate(
        (
            "导入前请删除示例行（第 2、3 行），否则会被一并导入。",
            "表头名称请勿修改；列的先后顺序可以调整，导入按表头识别。",
            "同一店铺内标题与内容完全相同的条目会被自动跳过，不会重复导入。",
            "兼容旧格式：含「一级分类 / 二级分类 / 话术标题 / 话术内容」的表格可直接导入。",
        )
    ):
        notes.cell(row=tail + offset, column=1, value=f"{offset + 1}.")
        notes.cell(row=tail + offset, column=2, value=line)

    workbook.save(filepath)


__all__ = [
    "COLUMN_CONTENT",
    "COLUMN_ENABLED",
    "COLUMN_TAGS",
    "COLUMN_TITLE",
    "STANDARD_HEADERS",
    "ParseResult",
    "export_workbook",
    "parse_enabled",
    "parse_rows",
    "parse_workbook",
    "write_template",
]
