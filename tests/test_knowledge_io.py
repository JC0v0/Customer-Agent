"""客服知识 Excel 导入 / 导出 / 模板回归。

重点锁三件事：
1. 导出的文件能被导入器原样读回（闭环），字段不丢
2. 旧的「一级分类 / 二级分类」表格继续可用
3. 坏数据被跳过且说得出原因，不会静默插入空标题
"""

from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from database.knowledge_service import KnowledgeService
from database.models import Base, Channel, Shop
from service import knowledge_io


def item(title="标题", content="内容", tags="物流", enabled=True):
    return SimpleNamespace(title=title, content=content, tags=tags, enabled=enabled)


class ParseRowsTests(unittest.TestCase):
    def test_standard_headers_in_any_order(self):
        result = knowledge_io.parse_rows(
            ["启用", "标签", "内容", "标题"],
            [["否", "物流,发货", "当天发货", "发货时间"]],
        )
        self.assertTrue(result.header_recognized)
        self.assertEqual(result.rows, [{
            "title": "发货时间", "content": "当天发货",
            "tags": "物流,发货", "enabled": False,
        }])

    def test_legacy_two_level_category_is_still_supported(self):
        result = knowledge_io.parse_rows(
            ["一级分类", "二级分类", "话术标题", "话术内容"],
            [["售后", "退换货", "七天无理由", "支持七天无理由"]],
        )
        self.assertTrue(result.header_recognized)
        self.assertEqual(result.rows[0]["tags"], "售后,退换货")
        self.assertTrue(result.rows[0]["enabled"])

    def test_unknown_headers_fall_back_to_legacy_positions_and_flag_it(self):
        result = knowledge_io.parse_rows(
            ["A", "B", "C", "D"],
            [["售后", "", "标题", "内容"]],
        )
        self.assertFalse(result.header_recognized)
        self.assertEqual(result.rows[0]["title"], "标题")
        self.assertEqual(result.rows[0]["tags"], "售后")

    def test_missing_required_fields_are_skipped_with_reason_and_row_number(self):
        result = knowledge_io.parse_rows(
            ["标题", "内容"],
            [["", "只有内容"], ["只有标题", ""], ["完整", "完整内容"]],
        )
        self.assertEqual(len(result.rows), 1)
        self.assertEqual(result.skipped[0][0], 2)
        self.assertIn("标题", result.skipped[0][1])
        self.assertEqual(result.skipped[1][0], 3)
        self.assertIn("内容", result.skipped[1][1])

    def test_blank_trailing_rows_are_ignored_silently(self):
        result = knowledge_io.parse_rows(
            ["标题", "内容"], [["有效", "内容"], ["", ""], [None, None]],
        )
        self.assertEqual(len(result.rows), 1)
        self.assertEqual(result.skipped, [])

    def test_unreadable_enabled_value_is_reported_not_guessed(self):
        result = knowledge_io.parse_rows(
            ["标题", "内容", "启用"], [["标题", "内容", "也许"]],
        )
        self.assertEqual(result.rows, [])
        self.assertIn("启用", result.skipped[0][1])

    def test_enabled_accepts_common_words_and_defaults_to_true(self):
        for raw, expected in (("是", True), ("否", False), ("TRUE", True),
                              ("0", False), ("", True), (None, True)):
            with self.subTest(raw=raw):
                self.assertIs(knowledge_io.parse_enabled(raw), expected)

    def test_pandas_nan_placeholders_are_treated_as_empty(self):
        result = knowledge_io.parse_rows(
            ["标题", "内容", "标签"], [["标题", "内容", "nan"]],
        )
        self.assertIsNone(result.rows[0]["tags"])

    def test_duplicate_tags_are_collapsed(self):
        result = knowledge_io.parse_rows(
            ["一级分类", "二级分类", "话术标题", "话术内容"],
            [["售后", "售后", "标题", "内容"]],
        )
        self.assertEqual(result.rows[0]["tags"], "售后")

    def test_describe_skipped_truncates_long_lists(self):
        rows = [["", f"内容{i}"] for i in range(60)]
        result = knowledge_io.parse_rows(["标题", "内容"], rows)
        text = result.describe_skipped(limit=10)
        self.assertEqual(len(text.splitlines()), 11)
        self.assertIn("另有 50 行", text)


class WorkbookRoundTripTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = str(Path(self.tmp.name) / "cs.xlsx")

    def test_exported_file_can_be_imported_back_without_loss(self):
        source = [
            item("发货时间", "当天 16:00 前发出", "物流,发货", True),
            item("七天无理由", "支持七天无理由退换", "售后", False),
            item("无标签", "这条没有标签", None, True),
        ]
        self.assertEqual(knowledge_io.export_workbook(self.path, source), 3)

        parsed = knowledge_io.parse_workbook(self.path)
        self.assertTrue(parsed.header_recognized)
        self.assertEqual(parsed.skipped, [])
        self.assertEqual(parsed.rows, [
            {"title": "发货时间", "content": "当天 16:00 前发出", "tags": "物流,发货", "enabled": True},
            {"title": "七天无理由", "content": "支持七天无理由退换", "tags": "售后", "enabled": False},
            {"title": "无标签", "content": "这条没有标签", "tags": None, "enabled": True},
        ])

    def test_formula_like_content_is_exported_as_plain_text(self):
        knowledge_io.export_workbook(self.path, [item(content="=1+1", tags="=SUM(A1)")])
        sheet = load_workbook(self.path).active
        self.assertEqual(sheet.cell(row=2, column=2).data_type, "s")
        self.assertEqual(sheet.cell(row=2, column=2).value, "=1+1")
        self.assertEqual(knowledge_io.parse_workbook(self.path).rows[0]["content"], "=1+1")

    def test_multiline_content_survives_round_trip(self):
        text = "第一行\n第二行\n第三行"
        knowledge_io.export_workbook(self.path, [item(content=text)])
        self.assertEqual(knowledge_io.parse_workbook(self.path).rows[0]["content"], text)

    def test_export_writes_standard_headers(self):
        knowledge_io.export_workbook(self.path, [item()])
        sheet = load_workbook(self.path).active
        headers = [sheet.cell(row=1, column=i).value for i in range(1, 5)]
        self.assertEqual(tuple(headers), knowledge_io.STANDARD_HEADERS)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = str(Path(self.tmp.name) / "template.xlsx")
        knowledge_io.write_template(self.path)

    def test_template_is_directly_importable(self):
        parsed = knowledge_io.parse_workbook(self.path)
        self.assertTrue(parsed.header_recognized)
        self.assertEqual(parsed.skipped, [])
        self.assertEqual(len(parsed.rows), 2)
        self.assertEqual(parsed.rows[0]["title"], "发货时间")

    def test_template_has_instruction_sheet_covering_every_column(self):
        workbook = load_workbook(self.path)
        self.assertIn("填写说明", workbook.sheetnames)
        text = "\n".join(
            str(cell.value) for row in workbook["填写说明"].iter_rows() for cell in row if cell.value
        )
        for column in knowledge_io.STANDARD_HEADERS:
            self.assertIn(column, text)
        self.assertIn("删除示例行", text)


class BatchImportPersistenceTests(unittest.TestCase):
    """导入落库：只用临时 SQLite，不触碰真实知识库。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        engine = create_engine(f"sqlite:///{Path(self.tmp.name) / 'kb.db'}")
        self.addCleanup(engine.dispose)
        Base.metadata.create_all(engine)
        sessions = sessionmaker(bind=engine)
        with mock.patch("database.knowledge_service.db_manager", SimpleNamespace(
            Session=sessions, engine=engine,
        )):
            self.service = KnowledgeService()
        with sessions() as session:
            session.add(Channel(id=1, channel_name="pinduoduo"))
            session.flush()
            session.add(Shop(id=1, channel_id=1, shop_id="s", shop_name="店铺"))
            session.commit()

    def _import(self, rows):
        return self.service.batch_import_customer_service(1, rows)

    def test_disabled_rows_are_stored_as_disabled(self):
        success, skipped = self._import([
            {"title": "启用条目", "content": "内容A", "tags": None, "enabled": True},
            {"title": "停用条目", "content": "内容B", "tags": None, "enabled": False},
        ])
        self.assertEqual((success, skipped), (2, 0))
        states = {
            cs.title: cs.enabled
            for cs in self.service.list_customer_service_with_disabled(1)
        }
        self.assertEqual(states, {"启用条目": True, "停用条目": False})

    def test_enabled_defaults_to_true_when_absent(self):
        self._import([{"title": "缺省", "content": "内容"}])
        self.assertTrue(self.service.list_customer_service_with_disabled(1)[0].enabled)

    def test_blank_title_or_content_never_reaches_the_database(self):
        success, skipped = self._import([
            {"title": "", "content": "只有内容"},
            {"title": "只有标题", "content": "   "},
            {"title": None, "content": None},
        ])
        self.assertEqual((success, skipped), (0, 3))
        self.assertEqual(self.service.list_customer_service_with_disabled(1), [])

    def test_duplicate_title_and_content_is_skipped(self):
        self._import([{"title": "重复", "content": "同样内容"}])
        success, skipped = self._import([{"title": "重复", "content": "同样内容"}])
        self.assertEqual((success, skipped), (0, 1))
        self.assertEqual(len(self.service.list_customer_service_with_disabled(1)), 1)

    def test_parsed_workbook_rows_flow_into_the_database(self):
        path = str(Path(self.tmp.name) / "cs.xlsx")
        knowledge_io.export_workbook(path, [
            SimpleNamespace(title="发货", content="当天发", tags="物流", enabled=True),
            SimpleNamespace(title="退换", content="七天无理由", tags="售后", enabled=False),
        ])
        success, skipped = self._import(knowledge_io.parse_workbook(path).rows)
        self.assertEqual((success, skipped), (2, 0))
        stored = {
            cs.title: (cs.tags, cs.enabled)
            for cs in self.service.list_customer_service_with_disabled(1)
        }
        self.assertEqual(stored, {"发货": ("物流", True), "退换": ("售后", False)})


if __name__ == "__main__":
    unittest.main()
